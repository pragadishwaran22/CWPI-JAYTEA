"""Read contractor sections from the two daily material-slip workbooks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import FilterColumn


SECTION_SHEETS = {
    "Production plan": "SHIFT {shift}",
    "Printing material": "PRINTING MATL REQ",
    "Packing material": "PACKING MATL REQ",
    "Speciality tea material": "SPECIALITY TEA MATL REQ",
    "Black tea material": "BLACK TEA MATL REQ",
}

PLAN_COLUMNS = [
    "Machine Line", "Item ID", "SKU Name", "Plan Qty", "UOM", "Machine No", "Blend Std",
]
MATERIAL_COLUMNS = [
    "Item ID", "Item Group", "Item Name", "UOM", "Required Qty", "Request Qty",
    "Available Floor Stock", "Pending Receive Requisition",
]
EXPORT_COLUMNS = [
    "Type", "Shift", "Date", "Item ID", "Item Name", "UOM",
    "Machine Line", "Machine No", "Plan Qty", "Blend Std", "Item Group",
    "Required Qty", "Request Qty", "Available Floor Stock",
    "Pending Receive Requisition",
]
@dataclass
class ClassifierShift:
    shift: str
    workbook_date: date | None
    sections: dict[str, pd.DataFrame]

    @property
    def contractors(self) -> set[str]:
        return {
            name
            for frame in self.sections.values()
            for name in frame["Contractor Name"].unique()
        }


def _text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip() if value is not None else ""


def _item_id(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _workbook_date(sheet) -> date | None:
    for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
        for index, value in enumerate(row[:-1]):
            if _text(value).upper().rstrip(":") == "DATE":
                raw = row[index + 1]
                if isinstance(raw, datetime):
                    return raw.date()
                if isinstance(raw, date):
                    return raw
                try:
                    return datetime.strptime(_text(raw), "%d-%m-%Y").date()
                except ValueError:
                    return None
    return None


def _read_section(sheet, *, production_plan: bool) -> pd.DataFrame:
    headers = PLAN_COLUMNS if production_plan else MATERIAL_COLUMNS
    records: list[dict[str, object]] = []
    contractor = ""
    machine_line = ""
    columns: dict[str, int] = {}
    unassigned_rows: list[int] = []

    aliases = {
        "ITEM ID": "Item ID",
        "SKU NAME": "SKU Name",
        "PLAN QTY": "Plan Qty",
        "MACHINE NO": "Machine No",
        "BLEND STD": "Blend Std",
        "ITEM GROUP": "Item Group",
        "ITEM NAME": "Item Name",
        "UOM": "UOM",
        "REQUIED QTY": "Required Qty",
        "REQUIRED QTY": "Required Qty",
        "REQUEST QTY": "Request Qty",
        "AVAILABLE FLOOR STOCK": "Available Floor Stock",
        "PENDING RECEIVE REQUISITION": "Pending Receive Requisition",
    }

    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        for value in row:
            heading = _text(value)
            contractor_match = re.match(r"^CONTRACTOR\s+NAME\s*:\s*(.+)$", heading, re.I)
            if contractor_match:
                contractor = _text(contractor_match.group(1))
                break
            if production_plan:
                line_match = re.match(r"^MACHINE\s+LINE\s*:\s*(.+)$", heading, re.I)
                if line_match:
                    machine_line = _text(line_match.group(1))
                    break
        else:
            normalized = [_text(value).upper() for value in row]
            if "ITEM ID" in normalized and ("SKU NAME" if production_plan else "ITEM NAME") in normalized:
                columns = {aliases[name]: i for i, name in enumerate(normalized) if name in aliases}
                continue
            name_column = "SKU Name" if production_plan else "Item Name"
            if "Item ID" not in columns or name_column not in columns:
                continue
            item_id = _item_id(row[columns["Item ID"]])
            item_name = _text(row[columns[name_column]])
            if not item_id or not item_name or item_id.upper() == "TOTAL":
                continue
            if not contractor:
                unassigned_rows.append(row_number)
                continue
            record: dict[str, object] = {"Contractor Name": contractor}
            if production_plan:
                record["Machine Line"] = machine_line
            for header in headers:
                if header == "Machine Line":
                    continue
                index = columns.get(header)
                record[header] = _item_id(row[index]) if header == "Item ID" and index is not None else row[index] if index is not None else None
            records.append(record)

    if unassigned_rows:
        raise ValueError(f"{sheet.title} has item rows without a contractor heading: {unassigned_rows[:10]}")
    return pd.DataFrame(records, columns=["Contractor Name", *headers])


def parse_classifier_workbook(source: bytes, expected_shift: str) -> ClassifierShift:
    """Read the five requested sheets without aggregating or changing quantities."""
    expected_shift = expected_shift.upper()
    if expected_shift not in {"A", "B"}:
        raise ValueError("Expected shift must be A or B.")
    workbook = load_workbook(BytesIO(source), read_only=True, data_only=True)
    try:
        lookup = {_text(name).upper(): name for name in workbook.sheetnames}
        required = {label: pattern.format(shift=expected_shift) for label, pattern in SECTION_SHEETS.items()}
        missing = [name for name in required.values() if name not in lookup]
        if missing:
            raise ValueError(f"Shift {expected_shift} workbook is missing: {', '.join(missing)}")
        plan = workbook[lookup[f"SHIFT {expected_shift}"]]
        sections = {
            label: _read_section(workbook[lookup[sheet_name]], production_plan=label == "Production plan")
            for label, sheet_name in required.items()
        }
        return ClassifierShift(expected_shift, _workbook_date(plan), sections)
    finally:
        workbook.close()


def contractor_sections(shift: ClassifierShift, contractor: str, search: str = "") -> dict[str, pd.DataFrame]:
    """Return source rows for one contractor, optionally matching text in any field."""
    result: dict[str, pd.DataFrame] = {}
    query = search.strip()
    for label, frame in shift.sections.items():
        selected = frame[frame["Contractor Name"].str.casefold() == contractor.casefold()].drop(columns="Contractor Name")
        if query:
            matches = selected.astype("string").fillna("").apply(
                lambda column: column.str.contains(query, case=False, regex=False)
            ).any(axis=1)
            selected = selected[matches]
        result[label] = selected.reset_index(drop=True)
    return result


def combined_contractor_sections(
    shifts: list[ClassifierShift], contractor: str, search: str = ""
) -> dict[str, pd.DataFrame]:
    """Show both shifts together within each source section, retaining source values."""
    per_shift = [(shift, contractor_sections(shift, contractor, search)) for shift in shifts]
    combined: dict[str, pd.DataFrame] = {}
    for label in SECTION_SHEETS:
        frames = []
        for shift, sections in per_shift:
            frame = sections[label].copy()
            frame.insert(0, "Shift", f"Shift {shift.shift}")
            frame.insert(1, "Date", shift.workbook_date.strftime("%d-%m-%Y") if shift.workbook_date else "")
            frames.append(frame)
        combined[label] = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    return combined


def _classifier_export_rows(
    shifts: list[ClassifierShift], contractor: str, search: str = ""
):
    """Map each source row into the shared export columns without changing values."""
    for label, frame in combined_contractor_sections(shifts, contractor, search).items():
        for source in frame.to_dict("records"):
            values = {header: source.get(header) for header in EXPORT_COLUMNS}
            values["Type"] = label
            values["Item Name"] = source.get("SKU Name") if label == "Production plan" else source.get("Item Name")
            yield tuple(values[header] for header in EXPORT_COLUMNS)


def create_classifier_excel_report(shifts: list[ClassifierShift], contractor: str, search: str = "") -> bytes:
    """Download one filterable sheet with independent production/material outlines."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contractor classifier"
    last_column = len(EXPORT_COLUMNS) + 1  # Narrow K divider keeps the two outlines separate.
    sheet.cell(1, 1, "CONTRACTOR")
    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_column)
    sheet.cell(1, 2, contractor)
    for row in sheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=last_column):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="F4D66A")
            cell.font = Font(bold=True, color="17191F", size=14 if cell.column == 2 else 10)
            cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, "Contractor-wise classifier report | Shift A and Shift B")
    sheet.cell(2, 1).font = Font(bold=True, color="F4D66A", size=12)
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor="27313E")
    sheet.row_dimensions[2].height = 24

    for col_number, header in enumerate(EXPORT_COLUMNS, start=1):
        display_column = col_number if col_number <= 10 else col_number + 1
        cell = sheet.cell(4, display_column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="27313E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(display_column)].width = min(max(len(header) + 4, 16), 38)
    sheet.cell(4, 11).fill = PatternFill("solid", fgColor="27313E")
    sheet.column_dimensions["K"].width = 3
    sheet.row_dimensions[4].height = 30

    for row_number, values in enumerate(_classifier_export_rows(shifts, contractor, search), start=5):
        for col_number, value in enumerate(values, start=1):
            display_column = col_number if col_number <= 10 else col_number + 1
            sheet.cell(row_number, display_column, value)
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.sheet_format.outlineLevelCol = 1
    for letter in ("G", "H", "I", "J", "L", "M", "N", "O", "P"):
        sheet.column_dimensions[letter].outlineLevel = 1
    sheet.freeze_panes = "D5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(last_column)}{max(4, sheet.max_row)}"
    sheet.auto_filter.filterColumn.append(FilterColumn(colId=10, hiddenButton=True))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
