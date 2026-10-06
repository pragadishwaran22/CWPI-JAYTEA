"""Regression checks for the all-contractor classifier workbook."""

from datetime import date
from io import BytesIO
import unittest

import pandas as pd
from openpyxl import load_workbook

from classifier_engine import (
    ClassifierShift,
    MATERIAL_COLUMNS,
    PLAN_COLUMNS,
    SECTION_SHEETS,
    create_all_classifier_excel_report,
    list_classifier_contractors,
)


def sample_shift(shift: str, contractors: list[str]) -> ClassifierShift:
    sections = {}
    for section_name in SECTION_SHEETS:
        columns = PLAN_COLUMNS if section_name == "Production plan" else MATERIAL_COLUMNS
        records = []
        for contractor in contractors:
            record = dict.fromkeys(columns)
            record["Contractor Name"] = contractor
            record["Item ID"] = f"{shift}-{contractor}-{section_name}"
            record["SKU Name" if section_name == "Production plan" else "Item Name"] = section_name
            records.append(record)
        sections[section_name] = pd.DataFrame(records, columns=["Contractor Name", *columns])
    return ClassifierShift(shift, date(2026, 10, 6), sections)


class ClassifierExportTests(unittest.TestCase):
    def test_all_contractors_get_separate_formatted_sheets(self):
        shifts = [sample_shift("A", ["BIPIN", "ANU"]), sample_shift("B", ["BIPIN", "ANU"])]
        workbook = load_workbook(BytesIO(create_all_classifier_excel_report(shifts)))
        self.assertEqual(workbook.sheetnames, ["ANU", "BIPIN"])
        for contractor in workbook.sheetnames:
            sheet = workbook[contractor]
            self.assertEqual(sheet["B1"].value, contractor)
            self.assertEqual(sheet["B1"].fill.fgColor.rgb[-6:], "F4D66A")
            self.assertEqual(sheet.freeze_panes, "D5")
            self.assertEqual(sheet.auto_filter.ref, "A4:P14")
            self.assertEqual(sheet["A4"].value, "Type")
            self.assertEqual(sheet["A5"].value, "Production plan")
            self.assertEqual(sheet["B5"].value, "Shift A")
            self.assertEqual(sheet["B10"].value, "Shift B")
            ids = [sheet.cell(row, 4).value for row in range(5, 15)]
            self.assertTrue(all(contractor in value for value in ids))

    def test_sheet_names_are_unique_and_excel_safe(self):
        names = ["Very Long Contractor Name/One A", "Very Long Contractor Name/One B"]
        shifts = [sample_shift("A", names), sample_shift("B", names)]
        workbook = load_workbook(BytesIO(create_all_classifier_excel_report(shifts)))
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertEqual(len({name.casefold() for name in workbook.sheetnames}), 2)
        self.assertTrue(all(len(name) <= 31 and "/" not in name for name in workbook.sheetnames))
        self.assertEqual({workbook[name]["B1"].value for name in workbook.sheetnames}, set(names))

    def test_contractor_names_ignore_case_only_duplicates(self):
        shifts = [sample_shift("A", ["BIPIN"]), sample_shift("B", ["Bipin"])]
        self.assertEqual(list_classifier_contractors(shifts), ["BIPIN"])
        workbook = load_workbook(BytesIO(create_all_classifier_excel_report(shifts)))
        self.assertEqual(len(workbook.sheetnames), 1)
        self.assertEqual(workbook.active.max_row, 14)


if __name__ == "__main__":
    unittest.main()
